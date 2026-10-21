"""Link saved survey geography evidence to the existing synthetic profile cohorts."""

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))
from ev_demand.geography import build_district_population_allocation as sources

require, resolve, rel, sha = sources.require, sources.resolve, sources.rel, sources.sha
CONTEXT_LABELS = {
    "T1": "At workplace, education place, shopping area, or other routine non-home place",
    "T2": "At home",
    "T3": "At tourism, leisure, hotel, restaurant, or destination place",
    "T4": "During a long trip / highway / en-route stop",
}
OWNER_QUESTIONS = {
    "home_district_answer": "4. What is your home district?",
    "away_geography_answer": "20. When you charge away from home, where does it normally happen?",
    "outside_district_answer": "21. If you often charge outside your home district, which district is most common?",
}
CONTEXT_PREFIXES = {
    "frequency_answer": "9. How often do you charge in each of the following situations?",
    "charger_answer": "13. What charger type do you normally use in each situation?",
    "start_window_answer": "14. At what time do you usually start charging in each situation?",
}


def text(value):
    return "" if value is None else " ".join(str(value).split())


def table(path, columns):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = reader.fieldnames or []
        require(len(fields) == len(set(fields)) and set(columns) <= set(fields),
                f"Missing or duplicate CSV columns: {rel(path)}")
        rows = list(reader)
    require(all(None not in row and all(v is not None for v in row.values()) for row in rows),
            f"Malformed CSV row: {rel(path)}")
    return rows


def build_evidence(owners, contexts, sheets, districts, bank_cohorts, profile_car_ids):
    owner_index = {row["owner_id"]: row for row in owners}
    require(len(owner_index) == len(owners) and owners, "Empty or duplicate owner feature records.")
    require(all(row["data_origin"] == "SYNTHETIC" for row in owners), "Use the existing synthetic feature snapshot.")
    context_index = {(row["owner_id"], row["context"]): row for row in contexts}
    require(len(context_index) == len(contexts)
            and set(context_index) == {(owner, code) for owner in owner_index for code in CONTEXT_LABELS},
            "Expected exactly four linked context records per retained owner.")
    require(set(bank_cohorts) <= set(owner_index) and profile_car_ids <= set(owner_index), "Unknown cohort owner ID.")
    require(profile_car_ids == {owner for owner, cohort in bank_cohorts.items()
                               if cohort == "common_energy_cases" and owner_index[owner]["vehicle_class_answer"] == "Car"},
            "The saved profile Car set differs from the common-bank Car set.")
    name_index = {text(row["district_name"]).casefold(): row for row in districts}
    require(len(name_index) == len(districts), "Duplicate canonical district names.")
    groups = {
        "all_retained": set(owner_index),
        "common_energy_cases": {owner for owner, cohort in bank_cohorts.items() if cohort == "common_energy_cases"},
        "common_Car_profiles": set(profile_car_ids),
        "open_upper_extension": {owner for owner, cohort in bank_cohorts.items() if cohort == "open_upper_extension"},
        "outside_screened_cohorts": set(owner_index) - set(bank_cohorts),
    }
    checks, evidence, context_columns = [], [], {}
    header_maps = {}
    for name, values in sheets.items():
        headers = [text(value) for value in values[0]]
        present = [value for value in headers if value]
        require(len(present) == len(set(present)), f"Duplicate survey headers: {name}")
        header_maps[name] = {value: column for column, value in enumerate(headers) if value}

    for owner_id, owner in owner_index.items():
        sheet_name, source_row = owner["source_sheet"], int(owner["source_row"])
        require(sheet_name in sheets and 2 <= source_row <= len(sheets[sheet_name]), f"Invalid survey source row: {owner_id}")
        values, headers = sheets[sheet_name][source_row - 1], header_maps[sheet_name]

        def raw(question):
            require(question in headers, f"Missing survey question: {question}")
            return text(values[headers[question]])

        def compare(record, field, question, code):
            actual, expected = text(record[field]), raw(question)
            require(actual == expected, f"Survey/feature answer mismatch: {owner_id}/{code}/{field}")
            checks.append({"owner_id": owner_id, "context": code, "feature_field": field,
                           "source_sheet": sheet_name, "source_row": source_row,
                           "source_column": headers[question] + 1, "question": question,
                           "source_answer": expected, "feature_answer": actual, "passed": True})
            return expected

        require(raw("Data_Origin").upper() == "SYNTHETIC", f"Unexpected raw source origin: {owner_id}")
        answers = {field: compare(owner, field, question, "OWNER") for field, question in OWNER_QUESTIONS.items()}
        for code, label in CONTEXT_LABELS.items():
            record = context_index[owner_id, code]
            require(record["data_origin"] == "SYNTHETIC" and record["source_sheet"] == sheet_name
                    and int(record["source_row"]) == source_row
                    and record["source_sha256"].lower() == owner["source_sha256"].lower(),
                    f"Context source-link mismatch: {owner_id}/{code}")
            for field, prefix in CONTEXT_PREFIXES.items():
                question = f"{prefix} [{label}]"
                compare(record, field, question, code)
                context_columns[sheet_name, code, field] = headers[question] + 1
        home = name_index.get(answers["home_district_answer"].casefold())
        outside = name_index.get(answers["outside_district_answer"].casefold())
        evidence.append({"owner_id": owner_id, "data_origin": "SYNTHETIC",
            "vehicle_class": owner["vehicle_class_answer"],
            "screened_cohort": bank_cohorts.get(owner_id, "outside_screened_cohorts"),
            "in_saved_Car_profiles": owner_id in profile_car_ids, **answers,
            "home_district_id": home["district_id"] if home else "",
            "home_match_status": "canonical_name_match" if home else "unmatched_for_review",
            "named_outside_district_id": outside["district_id"] if outside else "",
            "outside_match_status": "canonical_name_match" if outside else (
                "blank_answer" if not answers["outside_district_answer"] else "unmatched_text_for_review"),
            "named_outside_equals_home": (home["district_id"] == outside["district_id"]) if home and outside else "",
            "source_sheet": sheet_name, "source_row": source_row})
    definitions = []
    for sheet_name in sorted({row["source_sheet"] for row in owners}):
        count = sum(row["source_sheet"] == sheet_name for row in owners)
        for code, label in CONTEXT_LABELS.items():
            definitions.append({"source_sheet": sheet_name, "context": code, "survey_context": label,
                "linked_owner_records": count,
                **{f"{field}_column": context_columns[sheet_name, code, field] for field in CONTEXT_PREFIXES},
                "compared_answer_fields": count * 3, "source_answers_reproduced": True,
                "geographic_place_mapping_activated": False})
    summaries, counts, pairs = [], [], []
    for group, ids in groups.items():
        rows = [row for row in evidence if row["owner_id"] in ids]
        summaries.append({"population_group": group, "owners": len(rows),
            "home_district_matched": sum(row["home_match_status"] == "canonical_name_match" for row in rows),
            "away_geography_answer_present": sum(bool(row["away_geography_answer"]) for row in rows),
            "named_outside_district_matched": sum(row["outside_match_status"] == "canonical_name_match" for row in rows),
            "outside_answer_blank": sum(row["outside_match_status"] == "blank_answer" for row in rows),
            "outside_text_requires_review": sum(row["outside_match_status"] == "unmatched_text_for_review" for row in rows),
            "named_outside_equals_home": sum(row["named_outside_equals_home"] is True for row in rows)})
        for field in OWNER_QUESTIONS:
            frequencies = Counter(row[field] for row in rows)
            require(sum(frequencies.values()) == len(ids), "Geography answer denominator mismatch.")
            for answer, number in sorted(frequencies.items()):
                counts.append({"population_group": group, "question_key": field, "answer": answer,
                    "answer_is_blank": answer == "", "owners": number, "group_denominator": len(ids),
                    "owner_share": number / len(ids), "data_origin": "SYNTHETIC"})
        frequencies = Counter((row["home_district_id"], row["named_outside_district_id"])
                              for row in rows if row["home_district_id"] and row["named_outside_district_id"])
        for (home, outside), number in sorted(frequencies.items()):
            pairs.append({"population_group": group, "home_district_id": home,
                          "reported_most_common_outside_district_id": outside, "reporting_owners": number,
                          "group_denominator": len(ids), "statistic_basis": "owner_reports_not_session_flows"})
    return evidence, definitions, checks, summaries, counts, pairs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--features", required=True)
    parser.add_argument("--survey", required=True)
    parser.add_argument("--profiles", required=True)
    parser.add_argument("--output-root", default="data/processed/synthetic")
    args = parser.parse_args()
    watched = sources.Sources()
    watched.watch(Path(__file__).resolve())
    watched.watch(Path(sources.__file__).resolve())
    folder, survey_path, profile_folder = map(resolve, (args.features, args.survey, args.profiles))
    print("Verifying saved profile sources and the original synthetic survey workbook...")
    report = watched.package(profile_folder, "residence_attributed_profile_report.json")
    require(report.get("schema_version") == "residence_attributed_car_profiles_v1"
            and report.get("status") == "conditional_profiles_prepared" and report.get("behaviour_input_origin") == "SYNTHETIC"
            and report.get("checks") and all(value is True for value in report["checks"].values()), "Unsupported profile package.")
    owner_path = folder / "owners_features.csv"
    require(rel(owner_path) in report["input_sha256"], "Owner features are not bound to the saved profiles.")
    owners = table(owner_path, ("owner_id", "data_origin", "vehicle_class_answer", "source_file", "source_sha256",
                              "source_sheet", "source_row", *OWNER_QUESTIONS))
    require(owners, "No retained owners.")
    expected_hashes = {row["source_sha256"].strip().lower() for row in owners}
    require(len(expected_hashes) == 1 and re.fullmatch(r"[0-9a-f]{64}", next(iter(expected_hashes))),
            "Expected one shared raw survey snapshot with a valid SHA256.")
    watched.watch(survey_path, next(iter(expected_hashes)))
    candidates = []
    required_context = {"owner_id", "context", "frequency_answer", "charger_answer", "start_window_answer"}
    for path in sorted(folder.glob("*.csv")):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            columns = next(csv.reader(stream), [])
        if required_context <= set(columns):
            candidates.append(path)
    require(len(candidates) == 1, "Expected exactly one context feature CSV in the supplied feature folder.")
    context_path = candidates[0]
    watched.watch(context_path, report["input_sha256"].get(rel(context_path)))
    contexts = table(context_path, (*required_context, "data_origin", "source_sheet", "source_row", "source_sha256"))
    index_path = profile_folder / "district_profile_index.csv"
    require(index_path.name in report["output_sha256"], "District profile index is not hash-bound.")
    districts = {}
    for row in table(index_path, ("district_id", "district_name", "province_name")):
        identity = {key: row[key] for key in ("district_id", "district_name", "province_name")}
        require(row["district_id"] not in districts or districts[row["district_id"]] == identity, "District identity mismatch.")
        districts[row["district_id"]] = identity
    require(len(districts) == 25, "Expected 25 district identities.")
    bank_paths = [resolve(name) for name in report["input_sha256"] if Path(name).name == "screened_parameter_draws.csv"
                  and Path(name).parent.name.startswith("window_screened_parameters_")]
    require(len(bank_paths) == 1, "Missing or ambiguous source window-screened parameter bank.")
    bank_cohorts = {}
    for row in table(bank_paths[0], ("owner_id", "screen_cohort")):
        owner, cohort = row["owner_id"], row["screen_cohort"]
        require(cohort in {"common_energy_cases", "open_upper_extension"}
                and (owner not in bank_cohorts or bank_cohorts[owner] == cohort), "Inconsistent source cohort membership.")
        bank_cohorts[owner] = cohort
    review_path = resolve(report["source_review_reference"])
    require(rel(review_path) in report["input_sha256"], "Car review records are not hash-bound.")
    reviews = table(review_path, ("owner_id", "vehicle_class"))
    require(all(row["vehicle_class"] == "Car" for row in reviews), "Expected Car-only source review records.")
    profile_car_ids = {row["owner_id"] for row in reviews}
    require(len(profile_car_ids) == report["source_study_car_owners"], "Car profile owner denominator mismatch.")
    import openpyxl
    book = openpyxl.load_workbook(survey_path, read_only=True, data_only=False)
    try:
        sheet_names = {row["source_sheet"] for row in owners}
        require(sheet_names <= set(book.sheetnames), "Recorded survey sheet missing from the raw workbook.")
        sheets = {name: list(book[name].iter_rows(values_only=True)) for name in sheet_names}
    finally:
        book.close()
    evidence, definitions, checks, summaries, counts, pairs = build_evidence(
        owners, contexts, sheets, list(districts.values()), bank_cohorts, profile_car_ids)
    watched.unchanged()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    output = resolve(args.output_root) / f"charging_geography_evidence_{stamp}"
    output.mkdir(parents=True, exist_ok=False)
    exports = {
        "owner_geography_evidence.csv": (evidence, None), "context_definitions.csv": (definitions, None),
        "survey_answer_link_checks.csv": (checks, None), "geography_coverage_summary.csv": (summaries, None),
        "geography_answer_counts.csv": (counts, None),
        "reported_district_pairs.csv": (pairs, ("population_group", "home_district_id",
            "reported_most_common_outside_district_id", "reporting_owners", "group_denominator", "statistic_basis")),
    }
    for name, (rows, fields) in exports.items():
        with (output / name).open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    (output / "README.md").write_text(
        "# Synthetic charging geography evidence\n\n"
        f"All {len(owners)} retained owners remain represented. Group membership follows the saved parameter bank and Car profile records. "
        "The groups overlap: common_Car_profiles is a subset of common_energy_cases, and all_retained contains every group. "
        "Group counts must not be added as independent populations.\n\n"
        "The four T codes are charging contexts, verified against the frequency, charger and start-window questions in the original survey snapshot. "
        "They are not automatically the four geographic place typologies. Routine non-home includes education and other places; "
        "destination includes restaurants and leisure; en-route does not prove motorway charging; home does not prove a suburban location.\n\n"
        "Q20 records usual away-from-home geography. Q21 records a most-common outside district. They do not identify exact fractions of sessions, "
        "a destination separately for each context, or a session-level origin-destination matrix. District-pair rows count reporting owners only.\n\n"
        "Canonical district names are matched using case and whitespace normalization only. Blank answers and unmatched text remain distinct. "
        "Unmatched text can be a valid non-district response and is not automatically missing data or an exclusion reason. "
        "A named outside district equal to home is retained as a review flag. No values are imputed or corrected.\n\n"
        "These synthetic completeness cohorts do not establish national owner composition or travel behaviour. "
        "No routing probabilities, geographic place mappings, new eligibility rules, demand simulations, parameter draws or AI training are activated.\n",
        encoding="utf-8")
    watched.unchanged()
    result = {"schema_version": "charging_geography_evidence_v1", "status": "evidence_prepared",
        "data_origin": "SYNTHETIC", "owners_retained": len(owners), "context_records": len(contexts),
        "profile_Car_owners": len(profile_car_ids), "source_answer_comparisons": len(checks),
        "population_groups": summaries, "raw_survey_file": rel(survey_path), "raw_survey_sha256": next(iter(expected_hashes)),
        "context_feature_file": rel(context_path),
        "context_feature_was_already_profile_hash_bound": rel(context_path) in report["input_sha256"],
        "checks": {"recorded_source_hashes_verified": True, "source_answer_links_reproduced": True,
            "complete_owner_context_grid": True, "saved_Car_cohort_reproduced": True,
            "answer_count_denominators_checked": True, "source_files_unchanged": True},
        "context_specific_outside_shares_identified": False, "session_origin_destination_matrix_estimated": False,
        "geographic_place_mapping_activated": False, "eligibility_changed": False,
        "new_demand_simulations": 0, "new_parameter_draws": 0, "empirical_validation": False,
        "input_sha256": {rel(path): value for path, value in sorted(watched.hashes.items(), key=lambda item: str(item[0]))},
        "output_sha256": {path.name: sha(path) for path in sorted(output.iterdir()) if path.is_file()}}
    (output / "charging_geography_report.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print("Synthetic charging-geography evidence complete.")
    print(f"owners_retained: {len(owners)} | context_records: {len(contexts)} | source_answer_comparisons: {len(checks)}")
    for row in summaries:
        print(f"{row['population_group']}: owners={row['owners']}, home_matched={row['home_district_matched']}, "
              f"named_outside={row['named_outside_district_matched']}, outside_blank={row['outside_answer_blank']}, "
              f"other_outside_text={row['outside_text_requires_review']}")
    print("\nAway-from-home geography answers in the saved Car profile cohort:")
    for row in counts:
        if row["population_group"] == "common_Car_profiles" and row["question_key"] == "away_geography_answer":
            print(f"  {row['answer'] or '<blank>'}: {row['owners']}/{row['group_denominator']}")
    print("Source, context-definition, cohort and answer-denominator checks passed.")
    print(f"Output: {rel(output)}")
    print("Evidence only; session routing probabilities and geographic place mappings remain unestablished.")


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, KeyError, TypeError, csv.Error) as error:
        raise SystemExit(f"Charging geography evidence preparation failed: {error}") from error
