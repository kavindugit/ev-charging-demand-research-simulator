"""Extract DCS population and prepare two conditional Car allocation comparisons."""

import argparse
import csv
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
from fractions import Fraction
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
YEARS = tuple(range(2026, 2036))
SCENARIOS = ("low", "medium", "high")
FIELDS = ("population_total", "population_male", "population_female",
          "age_under_15", "age_15_59", "age_60_64", "age_65_plus")
COLUMNS = tuple("DEFGHIJ")
COUNT_ROWS = (12, *range(15, 64, 2))
PERCENT_ROWS = (13, *range(16, 65, 2))


def require(condition, message):
    if not condition:
        raise ValueError(message)


def resolve(value, base=ROOT):
    path = (base / value).resolve()
    require(path.is_relative_to(ROOT), f"Path outside repository: {path}")
    return path


def rel(path):
    return path.relative_to(ROOT).as_posix()


def sha(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_csv(path, required):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        require(set(required) <= set(reader.fieldnames or []), f"Missing CSV columns: {rel(path)}")
        return list(reader)


def english(value):
    return " ".join(re.findall(r"[A-Za-z]+", str(value or ""))).casefold()


def finite_decimal(value):
    result = Decimal(value)
    require(result.is_finite(), f"Non-finite value: {value}")
    return result


def decimal_text(value, places):
    """Round a Fraction exactly, with ties to even, for CSV display."""
    value = Fraction(value)
    scale = 10 ** places
    whole, remainder = divmod(abs(value.numerator) * scale, value.denominator)
    if 2 * remainder > value.denominator or (2 * remainder == value.denominator and whole % 2):
        whole += 1
    sign = "-" if value < 0 and whole else ""
    return f"{sign}{whole // scale}.{whole % scale:0{places}d}"


class Sources:
    def __init__(self):
        self.hashes = {}

    def watch(self, path, expected=None):
        require(path.is_file(), f"Source not found: {rel(path)}")
        if path not in self.hashes:
            self.hashes[path] = sha(path)
        if expected is not None:
            require(isinstance(expected, str) and re.fullmatch(r"[0-9a-fA-F]{64}", expected)
                    and self.hashes[path] == expected.lower(), f"Source hash mismatch: {rel(path)}")

    def package(self, folder, filename):
        report_path = folder / filename
        self.watch(report_path)
        report = read_json(report_path)
        for key in ("input_sha256", "output_sha256"):
            require(isinstance(report.get(key), dict) and report[key], f"Missing {key}: {rel(report_path)}")
        for name, expected in report["input_sha256"].items():
            self.watch(resolve(name), expected)
        for name, expected in report["output_sha256"].items():
            require(Path(name).name == name, f"Expected bare output filename: {name}")
            self.watch(resolve(name, folder), expected)
        return report

    def unchanged(self):
        for path, expected in self.hashes.items():
            require(sha(path) == expected, f"Source changed: {rel(path)}")


def district_identities(path):
    rows = read_csv(path, ("district_id", "district_name", "province_name", "year", "scenario"))
    identities, grid = {}, set()
    for row in rows:
        identity = {key: row[key].strip() for key in ("district_id", "district_name", "province_name")}
        key = identity["district_id"]
        require(re.fullmatch(r"DCS-[0-9]{2}", key) and all(identity.values()), "Invalid district identity.")
        require(key not in identities or identities[key] == identity, f"Inconsistent district: {key}")
        identities[key] = identity
        cell = (key, int(row["year"]), row["scenario"])
        require(cell not in grid, f"Duplicate district configuration: {cell}")
        grid.add(cell)
    require(len(identities) == 25 and len({r["province_name"] for r in identities.values()}) == 9,
            "Expected 25 districts and 9 provinces in the bound contract.")
    require(len({english(r["district_name"]) for r in identities.values()}) == 25, "Duplicate district names.")
    require(grid == {(key, year, case) for key in identities for year in YEARS for case in SCENARIOS},
            "District template must contain the complete 750-cell grid.")
    return list(identities.values())


def extract_population(path, identities):
    try:
        import openpyxl
    except ImportError as error:
        raise ValueError("openpyxl is required; install it in the project virtual environment.") from error
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        require(workbook.sheetnames == ["A1"], "Unexpected DCS sheet layout; inspect the workbook.")
        sheet = workbook["A1"]
        cells = list(sheet.iter_rows())
        require(len(cells) >= 64, "Incomplete DCS population table.")
        for address, expected in (("A3", "population by sex and age"), ("D4", "total number of persons"),
                                  ("E7", "male"), ("F7", "female"), ("G7", "less than years"),
                                  ("H7", "years"), ("I7", "years"), ("J7", "years and over")):
            require(expected in english(sheet[address].value), f"Unexpected DCS header at {address}.")
        require("2024" in str(sheet["A3"].value), "Expected the 2024 population table.")
        for address, ages in (("G7", (15,)), ("H7", (15, 59)), ("I7", (60, 64)), ("J7", (65,))):
            numbers = tuple(dict.fromkeys(int(v) for v in re.findall(r"[0-9]+", str(sheet[address].value))))
            require(numbers == ages, f"Unexpected age definition at {address}.")
        by_name = {english(row["district_name"]): row for row in identities}
        district_rows, national, audit, checks, seen = [], None, [], [], set()
        for index, row in enumerate(cells, start=1):
            populated = {cell.coordinate: cell.value for cell in row if cell.value is not None}
            label = row[0].value
            kind, name = "header" if index < 12 else "blank", ""
            if index in COUNT_ROWS:
                name = english(label)
                require(name == "sri lanka" or name in by_name, f"Unmatched district label at A{index}: {label}")
                require(name not in seen, f"Duplicate population region: {name}")
                seen.add(name)
                values = []
                for cell in row[3:10]:
                    value = cell.value
                    require(cell.data_type != "f" and not isinstance(value, bool)
                            and isinstance(value, (int, float)) and math.isfinite(value)
                            and value >= 0 and value == int(value), f"Invalid population count at {cell.coordinate}.")
                    values.append(int(value))
                counts = dict(zip(FIELDS, values))
                for topic, observed in (("sex_sum", values[1] + values[2]), ("age_sum", sum(values[3:]))):
                    require(observed == values[0], f"Population {topic} mismatch at source row {index}.")
                    checks.append({"region": name, "check": topic, "expected": values[0], "observed": observed, "passed": True})
                record = {**counts, "reference_year": 2024, "source_sheet": "A1", "source_row": index,
                          "source_label": label, "source_cells_json": json.dumps(dict(zip(FIELDS, (f"{c}{index}" for c in COLUMNS)))),
                          "data_origin": "PUBLIC_REFERENCE"}
                if name == "sri lanka":
                    require(national is None and index == 12, "Unexpected national population row.")
                    national, kind = {"region_name": "Sri Lanka", **record}, "national_count"
                else:
                    district_rows.append({**by_name[name], **record})
                    kind = "district_count"
            elif index in PERCENT_ROWS:
                require(label is None and row[3].value == 100, f"Unexpected percentage row: {index}")
                require(all(cell.data_type != "f" and not isinstance(cell.value, bool)
                            and isinstance(cell.value, (int, float)) and math.isfinite(cell.value)
                            and 0 <= cell.value <= 100 for cell in row[3:10]), f"Invalid percentage row: {index}")
                kind = "percentage_excluded"
            elif index >= 12:
                require(not populated, f"Unexpected populated row: {index}; inspect the DCS layout.")
            if index in COUNT_ROWS or index in PERCENT_ROWS:
                require(all(cell.value is None for cell in (*row[1:3], *row[10:])), f"Unexpected data columns at row {index}.")
            audit.append({"source_sheet": "A1", "source_row": index, "row_kind": kind,
                          "region": name, "cells_json": json.dumps(populated, ensure_ascii=False)})
        require(national is not None and len(district_rows) == 25 and seen == set(by_name) | {"sri lanka"},
                "National/district population coverage mismatch.")
        for field in FIELDS:
            total = sum(row[field] for row in district_rows)
            require(total == national[field], f"District-to-national population mismatch: {field}")
            checks.append({"region": "sri lanka", "check": f"district_sum/{field}",
                           "expected": national[field], "observed": total, "passed": True})
        require(national["population_total"] > 0, "National population must be positive.")
        return district_rows, national, audit, checks
    finally:
        workbook.close()


def scenario_counts(folder, report, reference):
    for name in ("scenario_design.json", "national_car_scenario_candidates.csv", "scenario_assumption_register.csv"):
        require(name in report["output_sha256"], f"Scenario file is not hash-bound: {name}")
    design = read_json(folder / "scenario_design.json")
    require(design.get("schema_version") == "conditional_car_scenario_design_v1"
            and design.get("status") == "draft_for_review" and design.get("vehicle_class_scope") == "Car"
            and design.get("count_representation") == "expected_vehicles"
            and design.get("count_origin") == "RESEARCH_SCENARIO_ASSUMPTION"
            and design.get("fleet_contract_activated") is False, "Expected an unactivated Car scenario draft.")
    require(design.get("years") == list(YEARS) and design.get("scenarios") == list(SCENARIOS)
            and design.get("interpolation") == "linear_share_between_2026_and_2035"
            and design.get("denominator_policy") == "fixed_2025_reference_not_projected_future_population"
            and design.get("count_decimal_places") == 6
            and design.get("public_population_reference") == reference, "Scenario design/reference mismatch.")
    population = finite_decimal(reference["reported_value"])
    require(reference["evidence_id"] == "DMT2025_CAR_POPULATION" and reference["reference_year"] == "2025"
            and reference["count_basis"] == "registered_population_all_fuels" and reference["vehicle_label"] == "MOTOR CARS"
            and reference["reported_fuel_category"] == "all_fuels" and reference["data_origin"] == "PUBLIC_REFERENCE"
            and population > 0 and population == population.to_integral_value(), "Invalid all-fuel Car reference.")
    rows = read_csv(folder / "national_car_scenario_candidates.csv", ("year", "scenario", "conditional_car_count",
        "proposed_electric_share_of_fixed_reference", "count_representation", "count_basis", "candidate_origin",
        "reference_all_fuel_car_population", "reference_population_year", "reference_evidence_id", "value_status",
        "share_assumption_id", "denominator_assumption_id", "path_assumption_id"))
    assumptions = read_csv(folder / "scenario_assumption_register.csv", ("assumption_id",))
    assumption_ids = {r["assumption_id"] for r in assumptions}
    require(len(assumption_ids) == len(assumptions) and "" not in assumption_ids, "Invalid scenario assumption IDs.")
    start = finite_decimal(design["start_share_2026"])
    ends = [finite_decimal(design["end_shares_2035"][case]) for case in SCENARIOS]
    require(Decimal(0) < start <= ends[0] < ends[1] < ends[2] <= Decimal(1), "Invalid scenario share ordering.")
    grid = set()
    for row in rows:
        key = (int(row["year"]), row["scenario"])
        require(key not in grid and key[0] in YEARS and key[1] in SCENARIOS, f"Invalid scenario cell: {key}")
        grid.add(key)
        with localcontext() as context:
            context.prec = 28
            share = start + (ends[SCENARIOS.index(key[1])] - start) * Decimal(key[0] - 2026) / Decimal(9)
            expected = (population * share).quantize(Decimal("0.000001"))
        require(finite_decimal(row["conditional_car_count"]) == expected
                and finite_decimal(row["proposed_electric_share_of_fixed_reference"]) == share
                and finite_decimal(row["reference_all_fuel_car_population"]) == population
                and row["reference_population_year"] == "2025" and row["reference_evidence_id"] == reference["evidence_id"]
                and row["count_representation"] == "expected_vehicles" and row["count_basis"] == "conditional_scenario_stock"
                and row["value_status"] == "draft_study_assumption"
                and row["candidate_origin"] == "RESEARCH_SCENARIO_ASSUMPTION", f"Scenario count/provenance mismatch: {key}")
        require(all(row[field] in assumption_ids for field in ("share_assumption_id", "denominator_assumption_id", "path_assumption_id")),
                f"Missing scenario assumption link: {key}")
    require(grid == {(year, case) for year in YEARS for case in SCENARIOS}, "Expected 30 national scenario cells.")
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True)
    parser.add_argument("--fleet-contract", required=True)
    parser.add_argument("--scenarios", required=True)
    parser.add_argument("--output-root", default="data/processed/synthetic")
    args = parser.parse_args()
    sources = Sources()
    sources.watch(Path(__file__).resolve())
    source, contract_folder, scenario_folder = map(resolve, (args.source, args.fleet_contract, args.scenarios))
    manifest_path = source / "source_manifest.json"
    sources.watch(manifest_path)
    manifest = read_json(manifest_path)
    require(manifest.get("reference_year") == 2024 and manifest.get("data_origin") == "PUBLIC_REFERENCE",
            "Expected the acquired 2024 DCS public-reference manifest.")
    workbook_path = resolve(manifest["source_file"])
    require(workbook_path.parent == source and workbook_path.suffix.lower() == ".xlsx", "Manifest source path mismatch.")
    sources.watch(workbook_path, manifest["source_sha256"])
    require(workbook_path.stat().st_size == manifest["source_bytes"], "Workbook byte count mismatch.")
    print("Verifying fleet-contract and national scenario source hashes...")
    contract = sources.package(contract_folder, "fleet_contract_report.json")
    require(contract.get("status") == "awaiting_fleet_inputs", "Expected the unfilled fleet contract.")
    for name in ("district_fleet_allocation_inputs_template.csv", "public_fleet_reference_counts.csv"):
        require(name in contract["output_sha256"], f"Contract file is not hash-bound: {name}")
    scenario_report = sources.package(scenario_folder, "scenario_design_report.json")
    require(scenario_report.get("schema_version") == "conditional_car_scenario_design_report_v1"
            and scenario_report.get("status") == "draft_for_review" and scenario_report.get("checks")
            and all(value is True for value in scenario_report["checks"].values()), "Scenario source checks are incomplete.")
    require(rel(contract_folder / "fleet_contract_report.json") in scenario_report["input_sha256"],
            "Scenario design is not bound to the supplied fleet contract.")
    identities = district_identities(contract_folder / "district_fleet_allocation_inputs_template.csv")
    references = read_csv(contract_folder / "public_fleet_reference_counts.csv", ("evidence_id",))
    reference = [r for r in references if r["evidence_id"] == "DMT2025_CAR_POPULATION"]
    require(len(reference) == 1, "Expected exactly one all-fuel Car population reference.")
    national_cases = scenario_counts(scenario_folder, scenario_report, reference[0])
    print("Reading DCS count rows; excluding the separate percentage rows...")
    population, national, audit, checks = extract_population(workbook_path, identities)
    for row in (*population, national):
        row.update(source_file=rel(workbook_path), source_sha256=manifest["source_sha256"])
    weights, counts, balances = [], [], []
    cases = (("equal_district_test", "ALLOC-EQUAL-01"), ("population_2024_test", "ALLOC-POP2024-01"))
    for case, assumption_id in cases:
        shares = {r["district_id"]: Fraction(1, 25) if case == "equal_district_test"
                  else Fraction(r["population_total"], national["population_total"]) for r in population}
        require(sum(shares.values(), Fraction(0)) == 1, f"Allocation shares do not sum to one: {case}")
        for row in population:
            identity = {key: row[key] for key in ("district_id", "district_name", "province_name")}
            share = shares[row["district_id"]]
            weights.append({"allocation_case": case, **identity, "allocation_share_display": decimal_text(share, 18),
                "share_numerator": share.numerator, "share_denominator": share.denominator,
                "population_reference_year": 2024, "district_population": row["population_total"],
                "national_population": national["population_total"], "allocation_basis": "conditional_owner_residence_proxy",
                "assumption_id": assumption_id, "horizon_assumption_id": "ALLOC-STATIC-01",
                "candidate_origin": "RESEARCH_SCENARIO_ASSUMPTION"})
        for scenario in national_cases:
            national_count = Fraction(finite_decimal(scenario["conditional_car_count"]))
            exact_total, displayed_total = Fraction(0), Decimal(0)
            for row in population:
                identity = {key: row[key] for key in ("district_id", "district_name", "province_name")}
                share = shares[row["district_id"]]
                exact = national_count * share
                displayed = decimal_text(exact, 6)
                exact_total += exact
                displayed_total += Decimal(displayed)
                counts.append({"allocation_case": case, **identity, "year": scenario["year"], "scenario": scenario["scenario"],
                    "vehicle_class_scope": "Car", "national_conditional_car_count": scenario["conditional_car_count"],
                    "share_numerator": share.numerator, "share_denominator": share.denominator,
                    "district_conditional_car_count": displayed, "count_representation": "expected_vehicles",
                    "count_basis": "conditional_scenario_stock", "allocation_basis": "conditional_owner_residence_proxy",
                    "allocation_assumption_id": assumption_id, "horizon_assumption_id": "ALLOC-STATIC-01",
                    "national_share_assumption_id": scenario["share_assumption_id"],
                    "national_denominator_assumption_id": scenario["denominator_assumption_id"],
                    "national_path_assumption_id": scenario["path_assumption_id"],
                    "value_status": "draft_study_assumption", "candidate_origin": "RESEARCH_SCENARIO_ASSUMPTION"})
            residual = displayed_total - Decimal(scenario["conditional_car_count"])
            require(exact_total == national_count and abs(residual) <= Decimal("0.0000125"), "District count balance failed.")
            balances.append({"allocation_case": case, "year": scenario["year"], "scenario": scenario["scenario"],
                "national_conditional_car_count": scenario["conditional_car_count"], "exact_fraction_balance_passed": True,
                "sum_exported_district_counts": str(displayed_total), "exported_minus_national": str(residual),
                "rounding_tolerance_vehicles": "0.0000125", "passed": True})
    assumptions = [
        {"assumption_id": "ALLOC-EQUAL-01", "rule": "Each district receives 1/25 of each national candidate count.",
         "rationale": "Equal-district benchmark; no claim about observed EV ownership.",
         "sensitivity_comparison": "Compare population_2024_test."},
        {"assumption_id": "ALLOC-POP2024-01", "rule": "Share = 2024 district population / 2024 national population.",
         "rationale": "Assumes equal EVs per resident, across all ages. Census residents are not observed EV owners or charging destinations.",
         "sensitivity_comparison": "Compare equal_district_test; later compare vehicle ownership or other justified evidence."},
        {"assumption_id": "ALLOC-STATIC-01", "rule": "Hold each allocation case's district shares fixed for 2026-2035.",
         "rationale": "Controlled comparison; no future population or ownership geography is forecast.",
         "sensitivity_comparison": "Review changing district shares when suitable evidence is available."},
    ]
    sources.unchanged()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    output = resolve(args.output_root) / f"district_car_allocation_candidates_{stamp}"
    output.mkdir(parents=True, exist_ok=False)
    tables = {"district_population_reference.csv": population, "national_population_reference.csv": [national],
              "population_source_row_audit.csv": audit, "population_reconciliation.csv": checks,
              "district_allocation_candidates.csv": weights, "district_car_count_candidates.csv": counts,
              "allocation_balance_checks.csv": balances, "allocation_assumptions.csv": assumptions}
    for name, rows in tables.items():
        with (output / name).open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    (output / "README.md").write_text(
        "# Conditional district Car allocation comparisons\n\n"
        "DCS population counts are public reference data. The two allocation rules and national Car share paths are study assumptions. "
        "Neither case measures district EV ownership. Each case contains 750 configurations; together they contain 1500 candidate rows.\n\n"
        "Population shares assume equal EVs per resident. Equal-district shares are a benchmark. Shares are fixed across 2026-2035. "
        "Counts describe a conditional owner-residence proxy, not charging destinations. Behaviour contexts and geographic place types are not allocated here.\n\n"
        "Exact rational shares conserve each supplied national count. CSV district counts are rounded to six decimal places; small signed residuals "
        "are retained in allocation_balance_checks.csv. No integer vehicles or new virtual owners are generated.\n\n"
        "No fleet contract or production weights are activated. No demand simulations, parameter draws, fitting or AI training occur. "
        "National scenarios remain drafts; Car registration-class and powertrain transfer assumptions remain unverified. "
        "The recorded acquisition hash checks snapshot consistency, not an independently published checksum.\n", encoding="utf-8")
    sources.unchanged()
    report = {"schema_version": "district_population_allocation_candidates_v1", "status": "draft_comparisons_prepared",
        "population_reference_year": 2024, "districts": 25, "provinces": 9,
        "national_population": national["population_total"], "percentage_rows_excluded": len(PERCENT_ROWS),
        "population_reconciliation_checks": len(checks), "allocation_cases": [c[0] for c in cases],
        "allocation_weight_rows": len(weights), "configurations_per_allocation_case": 750,
        "district_scenario_candidate_rows": len(counts), "allocation_balance_rows": len(balances),
        "vehicle_class_scope": "Car", "source_scenario_status": scenario_report["status"],
        "checks": {"recorded_source_hashes_verified": True, "district_identities_matched": True,
            "sex_age_and_national_totals_reconciled": True, "percentage_rows_excluded": True,
            "national_scenario_formulas_reproduced": True, "exact_fraction_allocation_conserved": True,
            "export_rounding_residuals_recorded": True, "source_files_unchanged": True},
        "district_EV_counts_observed": False, "charging_destination_allocation_established": False,
        "fleet_contract_activated": False, "production_weights_activated": False, "district_demand_ready": False,
        "new_demand_simulations": 0, "new_parameter_draws": 0, "new_virtual_owners": 0, "empirical_validation": False,
        "input_sha256": {rel(p): value for p, value in sorted(sources.hashes.items(), key=lambda item: str(item[0]))},
        "output_sha256": {p.name: sha(p) for p in sorted(output.iterdir()) if p.is_file()}}
    (output / "district_allocation_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("District population extraction and conditional allocation comparisons complete.")
    print(f"districts: 25 | provinces: 9 | population_2024: {national['population_total']}")
    print(f"percentage_rows_excluded: {len(PERCENT_ROWS)} | population_reconciliation_checks: {len(checks)}")
    print(f"allocation_cases: 2 | configurations_per_case: 750 | candidate_rows: {len(counts)}")
    print("Source, district, scenario-formula and exact allocation-balance checks passed.")
    print(f"Output: {rel(output)}")
    print("Conditional comparisons only; no fleet activation, new owners or demand simulations.")


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, KeyError, TypeError, InvalidOperation, csv.Error) as error:
        raise SystemExit(f"District population/allocation preparation failed: {error}") from error
